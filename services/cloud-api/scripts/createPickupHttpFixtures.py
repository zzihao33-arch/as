"""Create synthetic HTTP acceptance fixtures; requires pypdf, pillow, python-docx, openpyxl."""
from pathlib import Path
from pypdf import PdfWriter
from PIL import Image
from docx import Document
from openpyxl import Workbook

root = Path("node_modules/.cache/pickup-fixtures")
root.mkdir(parents=True, exist_ok=True)
writer = PdfWriter()
writer.add_blank_page(width=100, height=100)
writer.encrypt("synthetic-password")
with (root / "encrypted.pdf").open("wb") as output:
    writer.write(output)
Image.new("RGB", (64, 64), (35, 120, 160)).save(root / "synthetic.png")
document = Document()
document.add_paragraph("CM-HUB synthetic pickup acceptance")
document.save(root / "synthetic.docx")
workbook = Workbook()
workbook.active.append(["CM-HUB synthetic pickup acceptance", 1])
workbook.save(root / "synthetic.xlsx")
